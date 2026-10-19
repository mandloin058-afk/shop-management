"""Saaf Excel (shop_clean_data.xlsx) se Product, Purchase, Sale aur Expense import karta hai.

Pehle preview:   python manage.py import_excel shop_clean_data.xlsx
Asli import:     python manage.py import_excel shop_clean_data.xlsx --commit
"""
from decimal import Decimal
from datetime import datetime, date

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from shop.models import Customer, Expense, Product, Purchase, Sale


def rows(wb, name):
    ws = wb[name]
    head = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if any(v not in (None, "") for v in r):
            yield dict(zip(head, r))


def d(v):
    return v.date() if isinstance(v, datetime) else v


def dec(v):
    return Decimal(str(v)).quantize(Decimal("0.01"))


class Command(BaseCommand):
    help = "Saaf Excel se data import karo (--commit ke bina sirf preview)"

    def add_arguments(self, parser):
        parser.add_argument("path")
        parser.add_argument("--commit", action="store_true")

    def handle(self, *args, **o):
        wb = load_workbook(o["path"], data_only=True)
        prod = next(rows(wb, "Product"))
        P = list(rows(wb, "Purchases"))
        S = list(rows(wb, "Sales"))
        E = list(rows(wb, "Expenses"))

        buy_kg = sum(int(p["quantity_kg"]) for p in P)
        sell_kg = sum(int(s["quantity_kg"]) for s in S)
        if sell_kg > buy_kg:
            raise CommandError(f"Sales ({sell_kg} kg) purchases ({buy_kg} kg) se zyada hain, stock minus ho jayega.")
        for p in P:
            if not p["purchase_price_per_kg"]:
                raise CommandError(f"Purchase ka rate khali hai: {p}")

        self.stdout.write(f"Database mein abhi: Product {Product.objects.count()}, Purchase {Purchase.objects.count()}, "
                          f"Sale {Sale.objects.count()}, Expense {Expense.objects.count()}, Customer {Customer.objects.count()}")
        self.stdout.write(f"Import hoga: Purchases {len(P)} ({buy_kg/100} q), Sales {len(S)} ({sell_kg/100} q), "
                          f"Expenses {len(E)} (Rs {sum(int(e['amount']) for e in E)})")
        self.stdout.write(f"Import ke baad stock: {buy_kg - sell_kg} kg")
        if not o["commit"]:
            self.stdout.write(self.style.WARNING("Ye sirf preview tha. Asli import ke liye --commit lagao."))
            return

        with transaction.atomic():
            product, _ = Product.objects.get_or_create(
                name=prod["name"],
                defaults=dict(purchase_price=dec(prod["purchase_price_per_kg"]), selling_price=dec(prod["selling_price_per_kg"])),
            )
            product.purchase_price = dec(prod["purchase_price_per_kg"])  # sales ka cost isi se banta hai
            product.selling_price = dec(prod["selling_price_per_kg"])
            product.stock = 0
            product.save()
            for p in P:
                Purchase.objects.create(supplier=p["supplier"], product=product, quantity=int(p["quantity_kg"]),
                                        purchase_price=dec(p["purchase_price_per_kg"]), date=d(p["date"]))
            for s in S:
                cust, _ = Customer.objects.get_or_create(name=s["customer"])
                Sale.objects.create(customer=cust, product=product, quantity=int(s["quantity_kg"]),
                                    selling_price=dec(s["selling_price_per_kg"]),
                                    payment_method=s["payment_method"] or "Cash", date=d(s["date"]))
            for e in E:
                Expense.objects.create(category=e["category"], amount=dec(e["amount"]),
                                       description=e["description"], date=d(e["date"]))
            product.refresh_from_db()
        self.stdout.write(self.style.SUCCESS(f"Ho gaya. Makka ka stock ab {product.stock} kg hai."))
