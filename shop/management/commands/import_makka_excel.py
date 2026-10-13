"""
Purani Makka (corn) trading Excel file se data uthake shop_management
project ke Purchase/Sale tables me daalta hai.

IMPORTANT — is project ka schema simple hai (generic retail jaisa):
  - Purchase me sirf: supplier (naam), product, quantity, price, date
    hote hain. Koi "kata deduction" ya "partial payment" field nahi hai,
    aur koi notes field bhi nahi — isliye wo detail store nahi ho payegi,
    sirf final quantity aur rate store hoga (total khud calculate hota hai).
  - Sale me sirf ek "discount" field hai (alag se labour/gadi/mandi-tax
    nahi) — isliye Sheet2 ke saare kharche "discount" me jod diye jayenge.
  - Sheet2 ki wo rows jinme sirf kharcha likha hai par koi date/quantity
    nahi (jaise "Bag 9900"), unhe skip kiya jayega kyunki unhe kisi
    specific sale se jodne ka koi reliable tareeka nahi hai.

KAHA RAKHNA HAI:
    shop/management/commands/import_makka_excel.py
  (management aur commands folder nahi hai to pehle banao, dono me
   ek khali __init__.py file bhi honi chahiye)

KAISE CHALANA HAI:
    python manage.py import_makka_excel "path\\to\\Pappu_Uncle_1st_work_MakkaAutoRecovered.xlsx"
"""
import re
from datetime import datetime
from decimal import Decimal

from django.core.management.base import BaseCommand, CommandError

from shop.models import Product, Purchase, Sale


def D(value):
    """Plain float/int ko Decimal me convert karta hai (Django ke Decimal
    fields ke saath float ko seedha jodna/ghatana TypeError deta hai)."""
    return Decimal(str(round(value, 2)))

NUM = r"[\d,]+(?:\.\d+)?"
PRODUCT_NAME = "Makka (W)"


def to_float(s):
    return float(s.replace(",", ""))


def parse_quantity(text):
    """Returns (gross_kg, deduction_kg, party_name) from the messy
    quantity cell, e.g. '2910 kg Net -14 kg = 2896 Kg Total (Mukesh ji Patidar)'."""
    if not text:
        return None, 0.0, ""
    text = str(text)

    party_match = re.findall(r"\(([^)]+)\)", text)
    party_name = re.sub(r"\s+", " ", party_match[-1]).strip() if party_match else "Unknown Supplier"

    quintal_match = re.search(rf"({NUM})\s*Qu[uio]+ntal", text, re.IGNORECASE)
    if quintal_match:
        return to_float(quintal_match.group(1)) * 100, 0.0, party_name

    deduction = 0.0
    ded_match = re.search(rf"-\s*({NUM})\s*[Kk]g", text)
    if ded_match:
        deduction = to_float(ded_match.group(1))

    gross_kg = None
    for m in re.finditer(rf"({NUM})\s*[Kk]g", text):
        if "-" in text[max(0, m.start() - 2):m.start()]:
            continue
        gross_kg = to_float(m.group(1))
        break
    if gross_kg is None:
        any_num = re.search(NUM, text)
        gross_kg = to_float(any_num.group(0)) if any_num else 0.0

    return gross_kg, deduction, party_name


def parse_rate(text):
    if not text:
        return 0.0
    # allow an optional stray ")" between the number and "Rs" (seen in the source data)
    m = re.search(rf"({NUM})\s*\)?\s*Rs", str(text))
    return to_float(m.group(1)) if m else 0.0


def parse_money(text):
    """First plain number in a cell like '6800 Rs' or '32,040 Rs'."""
    if not text:
        return 0.0
    m = re.search(NUM, str(text))
    return to_float(m.group(0)) if m else 0.0


class Command(BaseCommand):
    help = "Import old Makka ledger Excel (Sheet1=purchases, Sheet2=sales) into shop_management."

    def add_arguments(self, parser):
        parser.add_argument("xlsx_path", type=str)

    def handle(self, *args, **options):
        try:
            import openpyxl
        except ImportError:
            raise CommandError("openpyxl install karo pehle: pip install openpyxl")

        path = options["xlsx_path"]
        wb = openpyxl.load_workbook(path, data_only=True)

        product, created = Product.objects.get_or_create(
            name=PRODUCT_NAME,
            defaults={"purchase_price": 17, "selling_price": 18, "stock": 0, "min_stock": 500},
        )
        if created:
            self.stdout.write(self.style.SUCCESS(f"Product '{PRODUCT_NAME}' bana diya."))

        # ---------------- Purchases (Sheet1) ----------------
        ws1 = wb["Sheet1"]
        purchases_created = 0
        for row in ws1.iter_rows(min_row=2, max_col=5, values_only=True):
            prod_name, date_val, qty_text, rate_text, _payment_text = row
            if not prod_name or not date_val:
                continue  # skip blank/junk rows (old sheet had a huge empty tail)

            gross_kg, deduction_kg, supplier_name = parse_quantity(qty_text)
            net_kg = max(0, round((gross_kg or 0) - (deduction_kg or 0)))
            rate = parse_rate(rate_text)
            date_obj = date_val.date() if isinstance(date_val, datetime) else date_val

            if net_kg <= 0 or rate <= 0:
                self.stdout.write(self.style.WARNING(f"Skip kiya (quantity/rate nahi mila): {qty_text!r}"))
                continue

            Purchase.objects.create(
                supplier=supplier_name,
                product=product,
                quantity=net_kg,
                purchase_price=D(rate),
                date=date_obj,
            )
            purchases_created += 1

        self.stdout.write(self.style.SUCCESS(f"{purchases_created} purchase entries import ho gaye."))

        # ---------------- Sales (Sheet2) ----------------
        ws2 = wb["Sheet2"]
        sales_created = 0
        for row in ws2.iter_rows(min_row=2, max_col=4, values_only=True):
            date_val, sell_text, labour_text, gadi_text = row
            if not date_val or not sell_text:
                continue  # rows with only a stray cost, no date/quantity — can't attach reliably

            quintal_match = re.search(rf"({NUM})\s*Qu[uio]+ntal", str(sell_text), re.IGNORECASE)
            if not quintal_match:
                continue
            qty_kg = round(to_float(quintal_match.group(1)) * 100)

            rate_match = re.search(rf"({NUM})\s*Rs/kg", str(sell_text))
            if rate_match:
                rate_per_kg = to_float(rate_match.group(1))
            else:
                rate_per_kg = product.selling_price
                self.stdout.write(self.style.WARNING(
                    f"Rate nahi mila is sale me ({sell_text!r}) — product ka default selling_price "
                    f"({product.selling_price}) use kiya hai, sale list me jaake sahi rate daal dena."
                ))

            extra_costs = parse_money(labour_text) + parse_money(gadi_text)
            date_obj = date_val.date() if isinstance(date_val, datetime) else date_val

            Sale.objects.create(
                customer=None,
                product=product,
                quantity=qty_kg,
                selling_price=D(rate_per_kg),
                discount=D(extra_costs),
                payment_method="Cash",
                date=date_obj,
            )
            sales_created += 1

        self.stdout.write(self.style.SUCCESS(f"{sales_created} sale entries import ho gaye."))
        self.stdout.write(self.style.WARNING(
            "Product ka stock ab Purchases/Sales ke hisaab se khud update ho gaya hoga. "
            "Admin panel (/admin/) me jaake Purchases aur Sales list check kar lena — "
            "kuch purani 'kata deduction' ya 'partial payment' details is simple schema me "
            "store nahi ho payi (upar comment me wajah likhi hai)."
        ))
